"""Build presentation with LESS content per slide so text stays large and readable."""

from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from lxml import etree

path = r"c:\Users\RavindraYadav\Documents\AI-Workbook-AWS\production-automation\aiops-production-automation-ansible-agentic\white_paper\Communication Sector Summit 2026_paper_presentation_submission_template.pptx"
prs = Presentation(path)

# Content area: below header, takes most of the slide
CONTENT_TOP = Emu(int(1.3 * 914400))
CONTENT_HEIGHT = Emu(int(5.9 * 914400))
CONTENT_LEFT = Emu(int(0.8 * 914400))
CONTENT_WIDTH = Emu(int(10.0 * 914400))

DARK = RGBColor(0x1A, 0x1A, 0x2E)
GRAY = RGBColor(0x55, 0x55, 0x55)
BODY = RGBColor(0x33, 0x33, 0x33)


def setup_slide(slide_idx):
    slide = prs.slides[slide_idx]
    shape = slide.shapes[1]
    shape.top = CONTENT_TOP
    shape.left = CONTENT_LEFT
    shape.width = CONTENT_WIDTH
    shape.height = CONTENT_HEIGHT
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    # Disable autofit so text stays at specified sizes
    txBody = tf._txBody
    bodyPr = txBody.find(qn("a:bodyPr"))
    if bodyPr is not None:
        for child in list(bodyPr):
            if child.tag in (qn("a:normAutofit"), qn("a:spAutoFit"), qn("a:noAutofit")):
                bodyPr.remove(child)
        etree.SubElement(bodyPr, qn("a:noAutofit"))
    return tf


def write(tf, lines):
    first = True
    for text, size, bold, italic, color in lines:
        if first:
            p = tf.paragraphs[0]
            first = False
        else:
            p = tf.add_paragraph()
        p.text = text
        if p.runs:
            r = p.runs[0]
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.italic = italic
            r.font.color.rgb = color


# ==========================================================
# SLIDE 1 — Problem (short, punchy, large font)
# Max ~10 lines so everything stays at 16-18pt
# ==========================================================
tf = setup_slide(0)
write(tf, [
    ("When Your Billing System Crashes at 3 AM, Who Fixes It?", 28, True, False, DARK),
    ("", 14, False, False, BODY),
    ("Energy & utilities IT teams manage billing, portals, meter data, and ERP.", 18, False, False, BODY),
    ("When something breaks, an engineer is paged \u2192 VPNs in \u2192 diagnoses \u2192 fixes.", 18, False, False, BODY),
    ("This takes 15\u201345 minutes. Every. Single. Time.", 18, True, False, DARK),
    ("", 14, False, False, BODY),
    ("\u2022  200\u2013400 alerts/day \u2014 spikes to 1,000+ during billing cycles", 16, False, False, BODY),
    ("\u2022  25\u201340 min MTTR for routine issues (restarts, disk, services)", 16, False, False, BODY),
    ("\u2022  23% of outages extended by incorrect manual fixes", 16, False, False, BODY),
    ("\u2022  Result: billing delays, portal downtime, missed SLAs", 16, False, False, BODY),
    ("", 14, False, False, BODY),
    ("What if IT infrastructure could fix routine problems itself?", 20, True, True, DARK),
])
print("Slide 1 done")

# ==========================================================
# SLIDE 2 — Solution (5 steps, large readable bullets)
# ==========================================================
tf = setup_slide(1)
write(tf, [
    ("AI-Powered Self-Healing IT Infrastructure", 28, True, False, DARK),
    ("Fixes problems before business users notice", 16, False, True, GRAY),
    ("", 12, False, False, BODY),
    ("1.  DETECT \u2014 CloudWatch catches CPU/memory/disk issues in seconds", 18, False, False, BODY),
    ("2.  UNDERSTAND \u2014 Enriches: which app, which team, which customers", 18, False, False, BODY),
    ("3.  DECIDE \u2014 Known problem \u2192 playbook. New \u2192 AI (Bedrock) analyzes", 18, False, False, BODY),
    ("4.  APPROVE \u2014 Operator gets email with root cause + Approve button", 18, False, False, BODY),
    ("5.  HEAL \u2014 Executes fix via AWS SSM. Validates. Reports back.", 18, False, False, BODY),
    ("", 12, False, False, BODY),
    ("Alert \u2192 Resolution in under 3 minutes", 20, True, False, DARK),
    ("", 12, False, False, BODY),
    ("Amazon Bedrock  |  Ansible + SSM  |  7-Layer Safety  |  ECS Fargate", 14, False, False, GRAY),
])
print("Slide 2 done")

# ==========================================================
# SLIDE 3 — Benefits (4 categories, 2 bullets each)
# ==========================================================
tf = setup_slide(2)
write(tf, [
    ("Why This Matters for Utility IT", 28, True, False, DARK),
    ("", 12, False, False, BODY),
    ("Business", 20, True, False, DARK),
    ("\u2022  80% MTTR reduction  |  Fewer SLA breaches with regulators", 16, False, False, BODY),
    ("\u2022  24/7 coverage without adding headcount", 16, False, False, BODY),
    ("", 10, False, False, BODY),
    ("Technical", 20, True, False, DARK),
    ("\u2022  AI learns from history \u2014 60\u201390% cheaper than always using top model", 16, False, False, BODY),
    ("\u2022  Zero SSH keys \u2014 all remediation via AWS Systems Manager", 16, False, False, BODY),
    ("", 10, False, False, BODY),
    ("Operational", 20, True, False, DARK),
    ("\u2022  Nothing executes without human approval  |  Full audit trail", 16, False, False, BODY),
    ("\u2022  Storm protection + circuit breakers auto-stop failing fixes", 16, False, False, BODY),
    ("", 10, False, False, BODY),
    ("Safety", 20, True, False, DARK),
    ("\u2022  Production DBs never auto-remediated  |  Blast radius classification", 16, False, False, BODY),
    ("\u2022  30-min approval expiry  |  Graceful fallback if AI unavailable", 16, False, False, BODY),
])
print("Slide 3 done")

# ==========================================================
# SLIDE 4 — Roadmap (compact, readable)
# ==========================================================
tf = setup_slide(3)
write(tf, [
    ("Roadmap & What\u2019s Next", 28, True, False, DARK),
    ("", 12, False, False, BODY),
    ("Today (Implemented)", 20, True, False, DARK),
    ("\u2713  Full detect \u2192 diagnose \u2192 approve \u2192 fix \u2192 verify pipeline", 16, False, False, BODY),
    ("\u2713  AI root cause (Bedrock) + human-in-the-loop email approval", 16, False, False, BODY),
    ("\u2713  7-layer safety guardrails  |  Production AWS (ECS Fargate)", 16, False, False, BODY),
    ("", 10, False, False, BODY),
    ("Next (Q3\u2013Q4 2026)", 20, True, False, DARK),
    ("\u2192  ServiceNow/BMC integration  |  Predictive alerting", 16, False, False, BODY),
    ("\u2192  AI-generated runbooks  |  Teams/Slack approval", 16, False, False, BODY),
    ("", 10, False, False, BODY),
    ("Vision", 20, True, False, DARK),
    ("10 engineers managing what used to need 40.", 18, False, True, BODY),
    ("IT shifts from firefighting \u2192 cloud migration & digital transformation.", 18, False, True, BODY),
    ("", 12, False, False, BODY),
    ("Ravindra Yadav  |  ravindya@in.ibm.com", 14, True, False, GRAY),
])
print("Slide 4 done")

output = r"c:\Users\RavindraYadav\Documents\AI-Workbook-AWS\production-automation\aiops-production-automation-ansible-agentic\white_paper\Energy_Utilities_IT_Final.pptx"
prs.save(output)
print(f"\nSaved: {output}")
