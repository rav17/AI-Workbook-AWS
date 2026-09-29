"""Add speaker notes to the PPTX using direct XML manipulation."""
from pptx import Presentation
from pptx.oxml.ns import qn
from lxml import etree

path = r"c:\Users\RavindraYadav\Documents\AI-Workbook-AWS\production-automation\aiops-production-automation-ansible-agentic\white_paper\Energy_Utilities_IT_Final.pptx"

prs = Presentation(path)

notes = [
    (
        "TALKING POINTS:\n\n"
        "Open with a story: 3:14 AM, billing cycle, meter data server 95% CPU. "
        "Engineer wakes, VPNs in, 20 min diagnosing routine issue fixed 12 times before.\n\n"
        "Key stats:\n"
        "- 200-400 alerts/day typical. During billing cycles spikes to 1000+\n"
        "- 25-40 min MTTR means billing stalls, customers locked out, SLA clocks tick\n"
        "- 23% of outages caused by the fix attempt itself (typos, wrong env)\n"
        "- 60%+ alerts are repetitive - same issue, same fix, every time\n"
        "- Alert fatigue: after 20+ alerts in shift, engineers start ignoring them\n\n"
        "Transition: What if infrastructure could fix routine problems itself, "
        "with human oversight but without human delay?"
    ),
    (
        "TALKING POINTS:\n\n"
        "Walk through each step - this is the core value proposition:\n\n"
        "1. DETECT: CloudWatch monitors every server. Alarm triggers in 60s.\n"
        "2. UNDERSTAND: Enrich with asset inventory - host, service, owner, impact.\n"
        "3. DECIDE: Known = proven playbook. New = Bedrock AI analyzes.\n"
        "   Model router: 70% Nova Lite ($0.0003), 20% Nova Pro ($0.001), "
        "10% Claude ($0.003). Saves 60-90% on AI costs.\n"
        "4. SAFETY: 7 guardrail checks short-circuit on DENY. Plan validator "
        "blocks dangerous AI commands (rm -rf, shutdown, iptables flush).\n"
        "5. APPROVE: HMAC-signed, DynamoDB-stored, 30-min expiry. "
        "Two-step (GET confirm, POST execute) prevents scanner auto-approval.\n"
        "6. HEAL: SSM RunCommand - no SSH keys. Stop conditions poll every 10s. "
        "Fleet breaker halts ALL automation if 5+ failures in 5 min.\n"
        "7. VERIFY: Baking period confirms alarm resolves. Outcome feeds back "
        "into model router for continuous improvement.\n\n"
        "Key message: Under 3 minutes for known problems with responsive operator."
    ),
    (
        "TALKING POINTS:\n\n"
        "Business Impact:\n"
        "- MTTR drops from 25 min to 3 min. Customers don't notice downtime.\n"
        "- 24/7 coverage: 3 AM billing crashes fixed in minutes.\n"
        "- Billing cycle protection: every minute of downtime delays thousands of invoices.\n\n"
        "Technical Excellence:\n"
        "- Model router: 70% alerts use Nova Lite (10x cheaper per call).\n"
        "- SSM: no bastion hosts, no SSH key rotation, no firewall rules.\n"
        "- Effectiveness scorer tracks what works. Auto-elevates if playbook starts failing.\n\n"
        "Operational Safety:\n"
        "- 7 guardrails with short-circuit. Fleet breaker: 5 failures = pause all.\n"
        "- Storm detection: 1000 alerts grouped to ~20 representative actions.\n"
        "- Inhibition: DatabaseDown active suppresses APITimeout, ConnectionRefused.\n\n"
        "Compliance:\n"
        "- Full audit trail: who approved, when, what executed, outcome.\n"
        "- DynamoDB conditional writes = exactly-once token consumption.\n"
        "- Structured logging with incident_id for post-incident reviews."
    ),
    (
        "TALKING POINTS:\n\n"
        "Architecture (high-level for IT leaders):\n"
        "- 6 CDK stacks: Network, Data, Compute, Observability, Monitoring, Simulation\n"
        "- ECS Fargate: serverless, no patching, auto-scales with alert volume\n"
        "- All traffic within VPC via endpoints (DynamoDB, S3, SQS, SSM, Bedrock)\n"
        "- Cost: ~$150/month at rest. Scales during billing peaks.\n\n"
        "Implemented (emphasize production-readiness):\n"
        "- NOT a POC. CI/CD with health gates, auto-rollback, vuln scanning.\n"
        "- 3-tier model routing + 4-model fallback. System never fully stops.\n"
        "- 12 YAML files: ops teams control without code deploys.\n"
        "- OpenTelemetry: end-to-end tracing from alert to remediation.\n\n"
        "Roadmap:\n"
        "- ServiceNow: automated change tickets, CMDB updates.\n"
        "- Predictive: if disk grows 2GB/day, alert BEFORE it fills.\n"
        "- AI-generated runbooks: observe manual fixes, auto-create playbook.\n\n"
        "Close strong: 10 engineers managing what used to need 40. "
        "IT shifts from firefighting to digital transformation."
    ),
]

for i, slide in enumerate(prs.slides):
    notes_slide = slide.notes_slide
    # Access the notes text body XML directly
    body = notes_slide.notes_text_frame._txBody

    # Remove all existing a:p elements
    for p_elem in body.findall(qn("a:p")):
        body.remove(p_elem)

    # Add new paragraphs for each line of the notes
    for line in notes[i].split("\n"):
        p_elem = etree.SubElement(body, qn("a:p"))
        r_elem = etree.SubElement(p_elem, qn("a:r"))
        t_elem = etree.SubElement(r_elem, qn("a:t"))
        t_elem.text = line

prs.save(path)

# Verify
prs2 = Presentation(path)
for i, s in enumerate(prs2.slides):
    n = s.notes_slide.notes_text_frame.text
    with open(f"white_paper/verify_s{i+1}.txt", "w", encoding="utf-8") as f:
        f.write(f"Content chars: {len(s.shapes[1].text_frame.text)}\n")
        f.write(f"Notes chars: {len(n)}\n")
        f.write(f"Notes preview: {n[:100]}\n")
